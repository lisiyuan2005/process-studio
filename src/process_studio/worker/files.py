"""Files the desktop's File menu reads and writes.

The flow as a table (Excel, CSV) or as a flow file (JSON, YAML); the
material, tool and recipe libraries as workbooks, in and out; a copy of the
whole workspace under another directory (Save as); and a way to show a
workspace folder in the system's file manager.
"""

from __future__ import annotations

import csv
import io
import json
import os
import shutil
import sqlite3
from contextlib import closing
import subprocess
import sys
from pathlib import Path
from typing import IO, Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..libraries import RecipeLibrary
from ..models import MaterialDefinition, ToolDefinition
from .errors import InvalidRequest, WorkspaceError
from .runner import build_document
from .workspace import load_project, open_repository

#: The columns a flow table can carry: an id the client asks for, the
#: heading it is written under, and the column width a workbook gives it.
#: A flow is exported for a reason -- a run sheet, a review, a tool list --
#: and which of these belong depends on the reason, so the client names the
#: ones it wants and gets them in this order.
FLOW_FIELDS: tuple[tuple[str, str, int], ...] = (
    ("index", "#", 4),
    ("name", "Step", 30),
    ("type", "Type", 12),
    ("tool", "Tool", 16),
    ("material", "Material", 12),
    ("mode", "Mode", 11),
    ("target", "Target (um)", 11),
    ("directional_fraction", "Directional fraction", 10),
    ("mask", "Mask", 18),
    ("keep", "Keep", 8),
    ("rates", "Rates (um/min)", 28),
    ("stops", "Stop layers", 16),
    ("other", "Other parameters (JSON)", 30),
    ("enabled", "Enabled", 8),
    ("status", "Status", 9),
    ("loop", "Loop", 14),
)
FLOW_FIELD_IDS = tuple(field for field, _heading, _width in FLOW_FIELDS)
FLOW_COLUMNS = [heading for _field, heading, _width in FLOW_FIELDS]
MATERIAL_COLUMNS = ["Material", "Category", "Color", "Opacity"]
#: "Recipes" is the machine's own recipe book, one cell, names separated
#: by semicolons -- a spreadsheet column per recipe would be unbounded.
TOOL_COLUMNS = ["Tool", "Group", "Recipes", "Notes"]
LIBRARY_KINDS = ("materials", "tools", "recipes")


def _cli():
    """The CLI's flow-file code, imported late: the CLI imports this worker."""
    from ..cli import flowfile
    from ..cli.session import Session

    return flowfile, Session


def _status_word(document: Mapping[str, Any], step_id: str) -> str:
    _, Session = _cli()
    return {"clean": "ready", "stale": "stale", "dirty": "not run"}.get(
        Session.status_of(document, step_id), "not run"
    )


def _chosen_fields(columns: Any) -> tuple[tuple[str, str, int], ...]:
    """The fields to write, in canonical order; None means all of them."""
    if columns is None:
        return FLOW_FIELDS
    if not isinstance(columns, (list, tuple)) or not all(isinstance(item, str) for item in columns):
        raise InvalidRequest("export_flow columns must be a list of column names.")
    wanted = set(columns)
    unknown = sorted(wanted - set(FLOW_FIELD_IDS))
    if unknown:
        raise InvalidRequest(
            f"export_flow does not have the column {unknown[0]!r}; it has "
            + ", ".join(FLOW_FIELD_IDS)
        )
    chosen = tuple(field for field in FLOW_FIELDS if field[0] in wanted)
    if not chosen:
        raise InvalidRequest("export_flow needs at least one column.")
    return chosen


def flow_rows(
    document: Mapping[str, Any], columns: Any = None, which: str = "simulation"
) -> list[list[Any]]:
    """One row per step, holding the chosen columns in canonical order.

    ``which`` picks what is written: the active branch, with the numbers
    the kernel was asked to build, or the experiment record -- what was done
    in the lab, each step with its tool and every setting. The record is a
    flow of its own, so its table has its own steps, not the branch's.
    """
    _, Session = _cli()
    fields = _chosen_fields(columns)
    if which not in ("simulation", "experiment"):
        raise InvalidRequest("export_flow values must be simulation or experiment.")
    steps = (
        list(document["project"].get("experimentSteps") or [])
        if which == "experiment"
        else Session.steps(document)
    )
    rows = []
    for index, step in enumerate(steps, start=1):
        parameters = dict(step.get("parameters", {}))
        parameters.pop("sketch_id", None)
        mode = parameters.pop("mode", None)
        target = parameters.pop("target", None)
        fraction = parameters.pop("directional_fraction", None)
        sketch_id = dict(step.get("parameters", {})).get("sketch_id")
        source = step.get("maskSource", "none")
        if source == "quick_sketch":
            mask = f"sketch:{sketch_id or 'default'}"
        elif source == "gds":
            mask = f"gds:{step.get('layer')}/{step.get('datatype') or 0}"
        else:
            mask = ""
        responses = step.get("materialResponses", {})
        rates = "; ".join(
            f"{name}={response['rateUmPerMin']}"
            for name, response in responses.items()
            if not response.get("stopLayer")
        )
        stops = "; ".join(name for name, response in responses.items() if response.get("stopLayer"))
        values = {
            "index": index,
            "name": step["name"],
            "type": step["processType"],
            "tool": step.get("tool", ""),
            "material": step.get("outputMaterial") or "",
            "mode": mode or "",
            "target": target,
            "directional_fraction": fraction,
            "mask": mask,
            "keep": step.get("keep", "inside") if mask else "",
            "rates": rates,
            "stops": stops,
            "other": json.dumps(parameters, ensure_ascii=False) if parameters else "",
            "enabled": bool(step.get("enabled", True)),
            # the record is never run
            "status": "" if which == "experiment" else _status_word(document, step["id"]),
            "loop": _loop_word(step),
        }
        rows.append([values[field] for field, _heading, _width in fields])
    return rows


def _loop_word(step: Mapping[str, Any]) -> str:
    loop = step.get("loop")
    if not loop:
        return ""
    return f"{loop.get('name') or 'Loop'} {int(loop.get('iteration', 0)) + 1}/{int(loop.get('repeat', 1))}"


def export_flow(
    root: Path,
    destination: Path,
    fmt: str,
    project_id: str | None = None,
    columns: Any = None,
    which: str = "simulation",
) -> dict[str, Any]:
    """Write the flow as a table (xlsx, csv) or a flow file (json, yaml).

    ``columns`` names the table columns to write (see ``FLOW_FIELDS``) and
    ``which`` the flow -- the active branch or the experiment record. The
    flow-file formats carry both, the record under ``experiment``, and
    ignore both: a flow file is what the workspace is rebuilt from.
    """
    flowfile, Session = _cli()
    repository = open_repository(root)
    document = build_document(root, repository, load_project(repository, project_id))
    fmt = fmt.lower()
    if fmt in ("json", "yaml", "yml"):
        flowfile.write_flow(
            destination.with_suffix(f".{fmt}") if destination.suffix == "" else destination,
            flowfile.flow_from_document(document),
        )
    elif fmt == "xlsx":
        fields = _chosen_fields(columns)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Process flow"
        sheet.append([heading for _field, heading, _width in fields])
        for row in flow_rows(document, columns, which):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, (_field, _heading, width) in enumerate(fields, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        workbook.save(destination)
    elif fmt == "csv":
        fields = _chosen_fields(columns)
        with destination.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow([heading for _field, heading, _width in fields])
            writer.writerows(flow_rows(document, columns, which))
    else:
        raise InvalidRequest("export_flow format must be xlsx, csv, json or yaml.")
    return {"path": str(destination), "steps": len(Session.steps(document)), "values": which}


def import_flow(root: Path, source: Path, output: IO[str]) -> dict[str, Any]:
    """Make the workspace match a flow file, the way ``flow apply`` does."""
    if not source.is_file():
        raise InvalidRequest("The selected flow file does not exist.")
    flowfile, Session = _cli()
    session = Session(root, events=output)
    flowfile.apply_flow(session, flowfile.read_flow(source))
    repository = open_repository(root)
    return {"path": str(source), "document": build_document(root, repository, load_project(repository))}


def export_library(root: Path, kind: str, destination: Path) -> dict[str, Any]:
    repository = open_repository(root)
    if kind == "recipes":
        RecipeLibrary(repository.load_recipes()).export_excel(destination)
        count = len(repository.load_recipes())
    elif kind == "materials":
        count = _write_sheet(
            destination, "Materials", MATERIAL_COLUMNS,
            [[m.name, m.category, m.color, m.opacity] for m in repository.load_materials()],
        )
    elif kind == "tools":
        count = _write_sheet(
            destination, "Tools", TOOL_COLUMNS,
            [[t.name, t.group, "; ".join(t.recipes), t.notes] for t in repository.load_tools()],
        )
    else:
        raise InvalidRequest(f"export_library kind must be one of {LIBRARY_KINDS}.")
    return {"path": str(destination), "count": count}


def import_library(root: Path, kind: str, source: Path, project_id: str | None = None) -> dict[str, Any]:
    """Add the rows of a workbook to a library; a row with a known name replaces it."""
    if not source.is_file():
        raise InvalidRequest("The selected workbook does not exist.")
    repository = open_repository(root)
    if kind == "recipes":
        try:
            library = RecipeLibrary.import_excel(source)
        except Exception as error:  # openpyxl raises workbook-specific errors
            raise InvalidRequest(f"Cannot read the recipe workbook: {error}") from error
        for recipe in library.recipes.values():
            repository.save_recipe(recipe)
        count = len(library.recipes)
    elif kind == "materials":
        existing = {m.name: m for m in repository.load_materials()}
        count = 0
        for row in _read_rows(source, MATERIAL_COLUMNS):
            name = str(row.get("Material") or "").strip()
            if not name:
                continue
            current = existing.get(name)
            material = MaterialDefinition(
                name,
                str(row.get("Category") or (current.category if current else "Other")),
                str(row.get("Color") or (current.color if current else "#7c83a0")),
                float(row.get("Opacity") if row.get("Opacity") not in (None, "") else (current.opacity if current else 1.0)),
                id=current.id if current else MaterialDefinition(name).id,
            )
            repository.save_material(material)
            count += 1
    elif kind == "tools":
        existing = {t.name: t for t in repository.load_tools()}
        count = 0
        for row in _read_rows(source, TOOL_COLUMNS):
            name = str(row.get("Tool") or "").strip()
            if not name:
                continue
            current = existing.get(name)
            recipes_cell = row.get("Recipes")
            tool = ToolDefinition(
                name,
                str(row.get("Group") or (current.group if current else "")),
                str(row.get("Notes") or (current.notes if current else "")),
                (
                    [part.strip() for part in str(recipes_cell).split(";")]
                    if recipes_cell
                    else list(current.recipes if current else [])
                ),
                id=current.id if current else ToolDefinition(name).id,
            )
            repository.save_tool(tool)
            count += 1
    else:
        raise InvalidRequest(f"import_library kind must be one of {LIBRARY_KINDS}.")
    return {"imported": count, "document": build_document(root, repository, load_project(repository, project_id))}


def _write_sheet(destination: Path, title: str, columns: list[str], rows: list[list[Any]]) -> int:
    if destination.suffix.lower() == ".csv":
        with destination.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            writer.writerows(rows)
        return len(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    workbook.save(destination)
    return len(rows)


def _read_rows(source: Path, columns: list[str]) -> list[dict[str, Any]]:
    """Rows of the first sheet (or a CSV) as dicts keyed by the header names."""
    if source.suffix.lower() == ".csv":
        with source.open("r", newline="", encoding="utf-8-sig") as handle:
            return [dict(row) for row in csv.DictReader(handle)]
    try:
        workbook = load_workbook(source, data_only=True)
    except Exception as error:
        raise InvalidRequest(f"Cannot read the workbook: {error}") from error
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell) if cell is not None else "" for cell in rows[0]]
    if not set(columns[:1]) <= set(headers):
        raise InvalidRequest(f"The sheet needs a {columns[0]!r} column; it has {headers}.")
    return [dict(zip(headers, values, strict=False)) for values in rows[1:]]


def copy_workspace(root: Path, destination: Path) -> dict[str, Any]:
    """Save as: the workspace copied under ``destination``, then opened.

    The database records each snapshot's absolute path, so the copy's
    rows are pointed at the copy's own snapshot directory.
    """
    root = root.resolve()
    destination = destination.expanduser().resolve()
    if destination == root or root in destination.parents:
        raise InvalidRequest("Save the workspace outside the current workspace directory.")
    if destination.exists() and any(destination.iterdir()):
        raise InvalidRequest(f"{destination} is not empty; choose a new or empty directory.")
    repository = open_repository(root)  # fails early on a non-workspace
    old_snapshots = repository.snapshot_directory.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    for item in root.iterdir():
        if item.name.startswith("_profile"):
            continue
        target = destination / item.name
        if item.is_dir():
            shutil.copytree(item, target, dirs_exist_ok=True)
        else:
            shutil.copy2(item, target)
    copied = open_repository(destination)
    new_snapshots = copied.snapshot_directory.resolve()
    with sqlite3.connect(copied.database_path) as connection:
        rows = connection.execute("SELECT id, path FROM snapshots").fetchall()
        for snapshot_id, path in rows:
            relative = Path(path).name
            connection.execute(
                "UPDATE snapshots SET path=? WHERE id=?", (str(new_snapshots / relative), snapshot_id)
            )
    del old_snapshots
    return {"root": str(destination), "document": build_document(destination, copied, load_project(copied))}


def reveal_path(path: Any) -> dict[str, Any]:
    """Show a directory (or a file's directory) in the system file manager."""
    if not isinstance(path, str) or not path:
        raise InvalidRequest("reveal_path requires a path.")
    target = Path(path).expanduser()
    if not target.exists():
        raise InvalidRequest(f"{target} does not exist.")
    folder = target if target.is_dir() else target.parent
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(folder))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(folder)])
        else:
            subprocess.Popen(["xdg-open", str(folder)])
    except OSError as error:
        raise WorkspaceError(f"Could not open {folder}: {error}") from error
    return {"path": str(folder)}


# -- a problem report ------------------------------------------------------------

#: A layout larger than this stays out of a report; the report says so.
REPORT_LAYOUT_LIMIT = 20 * 1024 * 1024


def write_report(
    root: Path,
    destination: Path,
    report: Any,
    log: Any = "",
) -> dict[str, Any]:
    """Pack what is needed to see a problem again into one zip.

    ``report`` is what the user said: a description, which flow they were
    looking at (``simulation`` or ``experiment``) and the ids of the steps
    they marked. The zip holds that, with the marked steps named and placed;
    the workspace as a flow file (both flows, materials, sketches, the
    window and the resolution), from which ``flow apply`` rebuilds it; the
    GDS layout when there is one; this build's version and platform; and the
    worker log the window had.
    """
    import platform
    import zipfile
    from datetime import datetime, timezone

    from .. import __version__

    if not isinstance(report, Mapping):
        raise InvalidRequest("write_report requires the report as an object.")
    which = str(report.get("flow") or "simulation")
    if which not in ("simulation", "experiment"):
        raise InvalidRequest("a report's flow must be simulation or experiment.")
    marked_ids = report.get("stepIds") or []
    if not isinstance(marked_ids, list) or not all(isinstance(item, str) for item in marked_ids):
        raise InvalidRequest("a report's stepIds must be a list of step ids.")

    flowfile, Session = _cli()
    repository = open_repository(root)
    project = load_project(repository)
    document = build_document(root, repository, project)
    branch = Session.branch(document)
    steps = list(document["project"].get("experimentSteps") or []) if which == "experiment" else Session.steps(document)
    statuses = document.get("stepStatuses", {}).get(branch["id"], {})
    wanted = set(marked_ids)
    marked = [
        {
            "index": index,
            "id": step["id"],
            "name": step["name"],
            "type": step["processType"],
            "loop": step.get("loop"),
            **({} if which == "experiment" else {"status": statuses.get(step["id"], "dirty")}),
        }
        for index, step in enumerate(steps, start=1)
        if step["id"] in wanted
    ]
    grid = document["project"]["grid"]
    summary = {
        "description": str(report.get("description") or ""),
        "flow": which,
        "branch": branch["name"],
        "markedSteps": marked,
        "created": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "app": {
            "version": __version__,
            "platform": platform.platform(),
            "python": platform.python_version(),
            "frozen": bool(getattr(sys, "frozen", False)),
        },
        "project": {
            "name": document["project"]["name"],
            "kernel": document["project"]["kernel"],
            "fidelity": document["project"].get("fidelity"),
            "resolutionUm": document["project"].get("resolutionUm"),
            "window": {key: grid[key] for key in ("xMin", "xMax", "yMin", "yMax", "zMin", "zMax")},
            "branches": [item["name"] for item in document["branches"]],
            "stepStatuses": {
                step["name"] if step["name"] else step["id"]: statuses.get(step["id"], "dirty")
                for step in Session.steps(document)
            },
        },
    }
    layout = Path(project.gds_path) if project.gds_path else None
    if layout is not None and layout.is_file() and layout.stat().st_size > REPORT_LAYOUT_LIMIT:
        summary["layoutLeftOut"] = f"{layout.name} is {layout.stat().st_size // (1024 * 1024)} MB"
        layout = None

    destination = destination if destination.suffix else destination.with_suffix(".zip")
    destination.parent.mkdir(parents=True, exist_ok=True)
    names = ["report.json", "flow.json", "log.txt"]
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr("report.json", json.dumps(summary, indent=2, ensure_ascii=False))
        bundle.writestr("flow.json", json.dumps(flowfile.flow_from_document(document), indent=2, ensure_ascii=False))
        bundle.writestr("log.txt", str(log or ""))
        if layout is not None and layout.is_file():
            bundle.write(layout, f"layout/{layout.name}")
            names.append(f"layout/{layout.name}")
    return {"path": str(destination), "files": names, "markedSteps": marked}


# -- a whole project as one zip ---------------------------------------------------

def _zip_folder(name: str) -> str:
    """The folder a project zip holds everything under: the project's name."""
    cleaned = "".join(ch if ch.isalnum() or ch in " -_." else "_" for ch in name).strip(" .")
    return cleaned or "Process Studio project"


def export_workspace_zip(root: Path, destination: Path, results: bool = False) -> dict[str, Any]:
    """The whole workspace as one zip, to send or to keep.

    Inside, one folder: the database (a consistent copy, taken with
    SQLite's backup, so a save going on cannot tear it), the sketches, the
    layout the project uses, and -- with ``results`` -- every step's stored
    result, so the copy opens with nothing to run. The 3D meshes kept
    beside results are left out (opening a 3D view builds them again), and
    so is anything nothing needs (old layouts, leftovers). Without results
    the copy's database forgets them too, so its steps open as not run
    rather than pointing at files that are not there.
    """
    import tempfile
    import zipfile

    from .workspace import DATABASE_NAME, SKETCH_DIRECTORY

    root = root.resolve()
    repository = open_repository(root)
    project = load_project(repository)
    destination = destination if destination.suffix else destination.with_suffix(".zip")
    if destination.resolve() == root or root in destination.resolve().parents:
        raise InvalidRequest("Save the zip outside the workspace directory.")
    folder = _zip_folder(project.name)
    snapshots = repository.snapshot_directory
    with tempfile.TemporaryDirectory(prefix="process-studio-export-") as scratch:
        copy = Path(scratch) / DATABASE_NAME
        # Every connection is closed by hand: `with` on one only commits,
        # and on Windows a database still open cannot be deleted, so the
        # scratch folder could not be cleaned up after the zip was written.
        with closing(sqlite3.connect(repository.database_path)) as source, closing(sqlite3.connect(copy)) as target:
            source.backup(target)
        files: list[tuple[Path, str]] = []
        layouts: dict[str, str] = {}
        with closing(sqlite3.connect(copy)) as connection, connection:
            connection.row_factory = sqlite3.Row
            # the layout each project uses, under layouts/ whatever its path was
            for row in connection.execute("SELECT id, gds_path FROM projects").fetchall():
                found = repository.on_disk(row["gds_path"], repository.layouts_directory)
                if found and Path(found).is_file():
                    name = Path(found).name
                    layouts[str(Path(found).resolve())] = name
                    connection.execute(
                        "UPDATE projects SET gds_path=? WHERE id=?", (f"layouts/{name}", row["id"])
                    )
                elif row["gds_path"]:
                    connection.execute("UPDATE projects SET gds_path=NULL WHERE id=?", (row["id"],))
            if results:
                rows = connection.execute("SELECT id, path FROM snapshots").fetchall()
                for row in rows:
                    found = repository.on_disk(row["path"], snapshots)
                    if found and Path(found).is_file():
                        name = Path(found).name
                        files.append((Path(found), f"{snapshots.name}/{name}"))
                        connection.execute(
                            "UPDATE snapshots SET path=? WHERE id=?", (f"{snapshots.name}/{name}", row["id"])
                        )
                    else:
                        connection.execute("DELETE FROM branch_snapshots WHERE snapshot_id=?", (row["id"],))
                        connection.execute("DELETE FROM snapshots WHERE id=?", (row["id"],))
            else:
                connection.execute("DELETE FROM branch_snapshots")
                connection.execute("DELETE FROM snapshots")
                exists = connection.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' AND name='worker_step_digests'"
                ).fetchone()
                if exists:
                    connection.execute("DELETE FROM worker_step_digests")
        with closing(sqlite3.connect(copy)) as connection:
            connection.execute("VACUUM")
        for source, name in layouts.items():
            files.append((Path(source), f"layouts/{name}"))
        sketches = root / SKETCH_DIRECTORY
        if sketches.is_dir():
            files += [(path, f"{SKETCH_DIRECTORY}/{path.name}") for path in sketches.iterdir() if path.is_file()]
        destination.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as bundle:
            bundle.write(copy, f"{folder}/{DATABASE_NAME}")
            for source, name in files:
                bundle.write(source, f"{folder}/{name}")
    return {
        "path": str(destination),
        "size": destination.stat().st_size,
        "results": results,
        "files": len(files) + 1,
    }


def import_workspace_zip(source: Path, destination: Path) -> dict[str, Any]:
    """Unpack a project zip into ``destination`` (a new or empty folder) and open it."""
    import zipfile

    from .workspace import DATABASE_NAME

    if not source.is_file():
        raise InvalidRequest("The selected project zip does not exist.")
    destination = destination.expanduser().resolve()
    if destination.exists() and any(destination.iterdir()):
        raise InvalidRequest(f"{destination} is not empty; choose a new or empty folder.")
    try:
        bundle = zipfile.ZipFile(source)
    except zipfile.BadZipFile as error:
        raise InvalidRequest(f"{source.name} is not a zip file.") from error
    with bundle:
        names = [name for name in bundle.namelist() if not name.endswith("/")]
        databases = [name for name in names if name.rsplit("/", 1)[-1] == DATABASE_NAME]
        if len(databases) != 1:
            raise InvalidRequest(f"{source.name} does not hold one Process Studio project.")
        prefix = databases[0][: -len(DATABASE_NAME)]
        wanted = [name for name in names if name.startswith(prefix)]
        for name in wanted:
            # nothing may land outside the folder it is unpacked into
            if destination not in (destination / name[len(prefix):]).resolve().parents:
                raise InvalidRequest(f"{source.name} holds an unsafe path: {name}")
        created = not destination.exists()
        destination.mkdir(parents=True, exist_ok=True)
        try:
            for name in wanted:
                target = destination / name[len(prefix):]
                target.parent.mkdir(parents=True, exist_ok=True)
                with bundle.open(name) as data, target.open("wb") as out:
                    shutil.copyfileobj(data, out)
            repository = open_repository(destination)
            document = build_document(destination, repository, load_project(repository))
        except Exception:
            # a half-unpacked project is no use to anyone
            if created:
                shutil.rmtree(destination, ignore_errors=True)
            raise
    return {"root": str(destination), "document": document}
