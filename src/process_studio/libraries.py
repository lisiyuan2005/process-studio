"""Material and recipe libraries with a deliberately simple Excel format."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import MaterialDefinition, MaterialResponse, ProcessType, Recipe


RECIPE_COLUMNS = [
    "Process Name",
    "Type",
    "Tool",
    "Recipe",
    "Material",
    "Time (min)",
    "Temperature (C)",
    "Target Thickness/Depth (um)",
    "Directional Fraction",
    "Rate (um/min)",
    "Stop Layer",
    "Extra Parameters (JSON)",
    # What the tool is set to when that is not what the kernel is asked to
    # build. Empty means the two are the same, which is the usual case and
    # which is why this is one column rather than a second set of them.
    "Experiment Parameters (JSON)",
    "Group",
]


class MaterialLibrary:
    def __init__(self, materials: list[MaterialDefinition] | None = None) -> None:
        self.materials = {material.name: material for material in materials or []}

    def add(self, material: MaterialDefinition) -> None:
        self.materials[material.name] = material

    def remove(self, name: str) -> None:
        self.materials.pop(name, None)


class RecipeLibrary:
    def __init__(self, recipes: list[Recipe] | None = None) -> None:
        self.recipes = {recipe.id: recipe for recipe in recipes or []}

    def add(self, recipe: Recipe) -> None:
        self.recipes[recipe.id] = recipe

    def export_excel(self, path: str | Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Recipes"
        sheet.append(RECIPE_COLUMNS)
        for recipe in self.recipes.values():
            responses = recipe.material_responses or {
                recipe.output_material or "": MaterialResponse(
                    recipe.output_material or "", float(recipe.parameters.get("rate", 0.0))
                )
            }
            for material, response in responses.items():
                common = {
                    key: value
                    for key, value in recipe.parameters.items()
                    if key
                    not in {
                        "time_min",
                        "temperature_c",
                        "target",
                        "directional_fraction",
                    }
                }
                sheet.append(
                    [
                        recipe.name,
                        recipe.process_type.value,
                        recipe.tool,
                        recipe.name,
                        material,
                        recipe.parameters.get("time_min"),
                        recipe.parameters.get("temperature_c"),
                        recipe.parameters.get("target"),
                        recipe.parameters.get("directional_fraction"),
                        response.rate_um_per_min,
                        response.stop_layer,
                        json.dumps(common, ensure_ascii=False),
                        (
                            ""
                            if recipe.experiment_parameters is None
                            else json.dumps(recipe.experiment_parameters, ensure_ascii=False)
                        ),
                        recipe.group,
                    ]
                )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(path)

    @classmethod
    def import_excel(cls, path: str | Path) -> "RecipeLibrary":
        workbook = load_workbook(path, data_only=True)
        sheet = workbook["Recipes"]
        headers = [cell.value for cell in sheet[1]]
        recipes_by_name: dict[str, Recipe] = {}
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values, strict=False))
            if not row.get("Recipe"):
                continue
            name = str(row["Recipe"])
            recipe = recipes_by_name.get(name)
            if recipe is None:
                extra = json.loads(row.get("Extra Parameters (JSON)") or "{}")
                parameters = dict(extra)
                mappings = {
                    "Time (min)": "time_min",
                    "Temperature (C)": "temperature_c",
                    "Target Thickness/Depth (um)": "target",
                    "Directional Fraction": "directional_fraction",
                }
                for column, key in mappings.items():
                    if row.get(column) is not None:
                        parameters[key] = row[column]
                experiment_text = str(row.get("Experiment Parameters (JSON)") or "").strip()
                recipe = Recipe(
                    name=name,
                    process_type=ProcessType(str(row["Type"])),
                    tool=str(row.get("Tool") or ""),
                    output_material=str(row.get("Material") or "") or None,
                    parameters=parameters,
                    experiment_parameters=json.loads(experiment_text) if experiment_text else None,
                    group=str(row.get("Group") or ""),
                )
                recipes_by_name[name] = recipe
            material = str(row.get("Material") or "")
            if material:
                recipe.material_responses[material] = MaterialResponse(
                    material=material,
                    rate_um_per_min=float(row.get("Rate (um/min)") or 0.0),
                    stop_layer=bool(row.get("Stop Layer") or False),
                )
        return cls(list(recipes_by_name.values()))
